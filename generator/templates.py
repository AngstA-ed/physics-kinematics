"""
Template functions that build each document type for the Curricula generator.

Each function accepts a python-docx Document and the relevant JSON data,
then populates the document with all required sections.
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from generator.styles import (
    set_document_defaults, add_title, add_heading_text, add_shaded_box,
    add_styled_table, add_bullet_list, add_lined_space, add_section_divider,
    add_styled_paragraph, style_run, set_cell_shading, set_paragraph_shading,
    DARK_BLUE, LIGHT_BLUE, LIGHT_GREEN, WHITE, BLACK, TABLE_HEADER_BG,
    FONT_NAME, TITLE_SIZE, HEADING_SIZE, BODY_SIZE, SMALL_SIZE, LIGHT_GRAY,
)


# ═══════════════════════════════════════════════════════════════════════════════
#  UNIT PLAN
# ═══════════════════════════════════════════════════════════════════════════════

def build_unit_plan(doc, course, unit):
    """
    Build a Unit Plan document.

    Sections:
    1. Unit title, course name, duration, PE codes
    2. Anchoring Phenomenon
    3. Essential Questions
    4. Unit Pacing/Sequence table
    5. Standards Map table
    6. CRSE Theme
    7. Assessment Overview
    8. Required Investigations
    """
    set_document_defaults(doc)

    # 1. Header block
    add_title(doc, f"{course} - Unit {unit.get('number', '')}: {unit.get('title', '')}")
    pe_codes = ", ".join(unit.get("performance_expectations", []))
    add_styled_paragraph(
        doc, f"Duration: {unit.get('days', 'TBD')} days  |  Performance Expectations: {pe_codes}",
        size=BODY_SIZE, italic=True, color=DARK_BLUE,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=Pt(8)
    )
    add_section_divider(doc)

    # 2. Anchoring Phenomenon
    add_heading_text(doc, "Anchoring Phenomenon")
    add_shaded_box(doc, "Phenomenon", unit.get("anchoring_phenomenon", "TBD"), LIGHT_BLUE)

    # 3. Essential Questions
    add_heading_text(doc, "Essential Questions")
    add_bullet_list(doc, unit.get("essential_questions", ["TBD"]))

    # 4. Unit Pacing / Sequence Table
    add_heading_text(doc, "Unit Pacing & Sequence")
    lessons = unit.get("lessons", [])
    pacing_rows = []
    for lesson in lessons:
        pacing_rows.append([
            str(lesson.get("number", "")),
            lesson.get("title", ""),
            str(lesson.get("days", "")),
            lesson.get("learning_target", "")[:120] + ("..." if len(lesson.get("learning_target", "")) > 120 else "")
        ])
    if pacing_rows:
        add_styled_table(
            doc,
            ["Lesson #", "Title", "Days", "Description"],
            pacing_rows,
            col_widths=[Inches(0.7), Inches(2.0), Inches(0.6), Inches(3.5)]
        )
    else:
        add_styled_paragraph(doc, "No lessons defined.", italic=True)

    doc.add_paragraph()  # spacer

    # 5. Standards Map Table
    add_heading_text(doc, "Standards Map")
    # Collect unique standards from all lessons
    all_pe, all_sep, all_ccc, all_dci = set(), set(), set(), set()
    for lesson in lessons:
        stds = lesson.get("standards", {})
        all_pe.update(stds.get("pe", []))
        all_sep.update(stds.get("sep", []))
        all_ccc.update(stds.get("ccc", []))
        all_dci.update(stds.get("dci", []))
    standards_rows = [[
        "\n".join(sorted(all_pe)) or "TBD",
        "\n".join(sorted(all_sep)) or "TBD",
        "\n".join(sorted(all_ccc)) or "TBD",
        "\n".join(sorted(all_dci)) or "TBD",
    ]]
    add_styled_table(
        doc,
        ["PE", "SEP", "CCC", "DCI"],
        standards_rows,
        col_widths=[Inches(1.2), Inches(2.0), Inches(1.5), Inches(2.5)]
    )

    doc.add_paragraph()

    # 6. CRSE Theme
    add_heading_text(doc, "CRSE Theme")
    add_shaded_box(doc, "Theme", unit.get("crse_theme", "TBD"), LIGHT_GREEN)

    # 7. Assessment Overview
    add_heading_text(doc, "Assessment Overview")
    assess = unit.get("assessment_overview", {})
    formative = assess.get("formative", [])
    summative = assess.get("summative", [])
    if formative:
        add_styled_paragraph(doc, "Formative Assessments:", bold=True, space_before=Pt(4))
        add_bullet_list(doc, formative)
    if summative:
        add_styled_paragraph(doc, "Summative Assessments:", bold=True, space_before=Pt(4))
        add_bullet_list(doc, summative)
    if not formative and not summative:
        add_styled_paragraph(doc, "TBD", italic=True)

    # 8. Required Investigations
    investigations = unit.get("investigations", [])
    if investigations:
        add_heading_text(doc, "Required Investigations")
        add_bullet_list(doc, investigations)

    return doc


# ═══════════════════════════════════════════════════════════════════════════════
#  LESSON PLAN
# ═══════════════════════════════════════════════════════════════════════════════

def build_lesson_plan(doc, course, unit, lesson):
    """
    Build a Lesson Plan document.

    Sections:
    1. Header box (Unit name, Lesson #, Title, PE codes)
    2. Opening Circle (light blue)
    3. Do Now (light green)
    4. Learning Target
    5. Standards Alignment table
    6. Materials & Preparation
    7. 5E Instructional Sequence
    8. Differentiation Box
    9. CRSE Connections
    10. Hochman Integration
    """
    set_document_defaults(doc)

    # 1. Header box
    unit_title = f"Unit {unit.get('number', '')}: {unit.get('title', '')}"
    lesson_title = f"Lesson {lesson.get('number', '')}: {lesson.get('title', '')}"
    pe_codes = ", ".join(lesson.get("standards", {}).get("pe", unit.get("performance_expectations", [])))

    p = doc.add_paragraph()
    set_paragraph_shading(p, TABLE_HEADER_BG)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"{course}\n")
    style_run(run, size=TITLE_SIZE, bold=True, color=WHITE)
    run2 = p.add_run(f"{unit_title}\n{lesson_title}\n")
    style_run(run2, size=HEADING_SIZE, bold=True, color=WHITE)
    run3 = p.add_run(f"PE: {pe_codes}")
    style_run(run3, size=BODY_SIZE, italic=True, color=RGBColor(0xD6, 0xE4, 0xF0))

    # 2. Opening Circle (light blue)
    add_heading_text(doc, "Suggested Opening Circle (2-3 min)")
    add_shaded_box(doc, "Suggested Community Prompt", lesson.get("circle_prompt", "TBD"), LIGHT_BLUE)

    # 3. Do Now (light green)
    add_heading_text(doc, "Do Now (3-5 min)")
    add_shaded_box(doc, "Activity", lesson.get("do_now", "TBD"), LIGHT_GREEN)

    # 4. Learning Target
    add_section_divider(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("Learning Target: ")
    style_run(run, size=HEADING_SIZE, bold=True, color=DARK_BLUE)
    run2 = p.add_run(lesson.get("learning_target", "TBD"))
    style_run(run2, size=HEADING_SIZE, bold=True)
    add_section_divider(doc)

    # 5. Standards Alignment Table
    add_heading_text(doc, "Standards Alignment")
    stds = lesson.get("standards", {})
    stds_rows = [[
        "\n".join(stds.get("pe", ["TBD"])),
        "\n".join(stds.get("sep", ["TBD"])),
        "\n".join(stds.get("ccc", ["TBD"])),
        "\n".join(stds.get("dci", ["TBD"])),
    ]]
    add_styled_table(
        doc,
        ["PE", "SEP", "CCC", "DCI"],
        stds_rows,
        col_widths=[Inches(1.2), Inches(2.0), Inches(1.5), Inches(2.5)]
    )
    doc.add_paragraph()

    # 6. Materials & Preparation
    materials = lesson.get("materials", [])
    if materials:
        add_heading_text(doc, "Materials & Preparation")
        add_bullet_list(doc, materials)

    # 7. 5E Instructional Sequence
    add_heading_text(doc, "5E Instructional Sequence")
    five_e = lesson.get("five_e", {})
    phase_labels = [
        ("engage", "Engage"),
        ("explore", "Explore"),
        ("explain", "Explain"),
        ("elaborate", "Elaborate"),
        ("evaluate", "Evaluate"),
    ]
    for key, label in phase_labels:
        phase = five_e.get(key, {})
        time_est = phase.get("time", "TBD")
        desc = phase.get("description", "TBD")
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run_label = p.add_run(f"{label} ({time_est}): ")
        style_run(run_label, size=BODY_SIZE, bold=True, color=DARK_BLUE)
        run_desc = p.add_run(desc)
        style_run(run_desc, size=BODY_SIZE)

    doc.add_paragraph()

    # 8. Differentiation Box
    add_heading_text(doc, "Differentiation")
    diff = lesson.get("differentiation", {})
    diff_labels = [
        ("approaching", "Approaching"),
        ("on_level", "On-Level"),
        ("advanced", "Advanced"),
        ("ell", "ELL"),
        ("iep", "IEP"),
    ]
    diff_rows = [[label, diff.get(key, "TBD")] for key, label in diff_labels]
    add_styled_table(
        doc,
        ["Level", "Accommodations / Modifications"],
        diff_rows,
        col_widths=[Inches(1.5), Inches(5.5)]
    )
    doc.add_paragraph()

    # 9. CRSE Connections
    add_heading_text(doc, "CRSE Connections")
    add_shaded_box(
        doc, "Culturally Responsive-Sustaining Education",
        lesson.get("crse_connection", "TBD"),
        LIGHT_GREEN
    )

    # 10. Hochman Integration
    add_heading_text(doc, "Suggested Hochman Writing Activity")
    add_shaded_box(
        doc, "Suggested Literacy Activity",
        lesson.get("hochman_activity", "TBD"),
        LIGHT_BLUE
    )

    return doc


# ═══════════════════════════════════════════════════════════════════════════════
#  STUDENT NOTES
# ═══════════════════════════════════════════════════════════════════════════════

def build_student_notes(doc, course, unit, lesson):
    """
    Build a Student Notes document.

    Sections:
    1. Header with unit/lesson info
    2. Learning Target
    3. Key Vocabulary (term + definition with blanks)
    4. Guided Notes with fill-in-the-blank
    5. Summary with Because/But/So sentence frame
    """
    set_document_defaults(doc)

    # 1. Header
    add_title(doc, f"{course} - Student Notes")
    add_styled_paragraph(
        doc,
        f"Unit {unit.get('number', '')}: {unit.get('title', '')}  |  "
        f"Lesson {lesson.get('number', '')}: {lesson.get('title', '')}",
        size=BODY_SIZE, italic=True, color=DARK_BLUE,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=Pt(4)
    )
    add_styled_paragraph(
        doc, f"Name: ________________________  Date: ____________  Period: _____",
        size=BODY_SIZE, alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=Pt(4), space_after=Pt(8)
    )
    add_section_divider(doc)

    # 2. Learning Target
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    set_paragraph_shading(p, LIGHT_BLUE)
    run = p.add_run("Learning Target: ")
    style_run(run, size=BODY_SIZE, bold=True, color=DARK_BLUE)
    run2 = p.add_run(lesson.get("learning_target", "TBD"))
    style_run(run2, size=BODY_SIZE)

    # 3. Key Vocabulary
    vocab = lesson.get("vocabulary", [])
    if vocab:
        add_heading_text(doc, "Key Vocabulary")
        vocab_rows = []
        for v in vocab:
            term = v.get("term", "")
            definition = v.get("definition", "")
            # Create a fill-in-the-blank version: replace some words with blanks
            blank_def = _make_fill_in_blank(definition)
            vocab_rows.append([term, blank_def])
        add_styled_table(
            doc,
            ["Term", "Definition (fill in the blanks)"],
            vocab_rows,
            col_widths=[Inches(2.0), Inches(5.0)]
        )
        doc.add_paragraph()

    # 4. Guided Notes
    guided = lesson.get("guided_notes_outline", [])
    if guided:
        add_heading_text(doc, "Guided Notes")
        for section in guided:
            add_styled_paragraph(doc, section, bold=True, size=BODY_SIZE, space_before=Pt(6))
            add_lined_space(doc, num_lines=3)
    else:
        add_heading_text(doc, "Guided Notes")
        add_lined_space(doc, num_lines=6)

    # 5. Summary - Because/But/So
    add_section_divider(doc)
    add_heading_text(doc, "Summary")
    topic = lesson.get("title", "today's lesson")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    set_paragraph_shading(p, LIGHT_GREEN)
    run = p.add_run("Complete the Because/But/So sentence frame about ")
    style_run(run, size=BODY_SIZE, italic=True)
    run2 = p.add_run(f"{topic}:")
    style_run(run2, size=BODY_SIZE, bold=True, italic=True)

    frames = [
        f"{topic} is important because",
        f"Scientists study {topic.lower()}, but",
        f"Understanding {topic.lower()} helps us, so",
    ]
    for frame in frames:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"{frame} ")
        style_run(run, size=BODY_SIZE, bold=True)
        run2 = p.add_run("_" * 55)
        style_run(run2, size=BODY_SIZE, color=RGBColor(0xBB, 0xBB, 0xBB))

    return doc


def _make_fill_in_blank(definition):
    """Replace ~every third word with a blank for fill-in-the-blank style."""
    words = definition.split()
    if len(words) <= 3:
        return definition
    result = []
    for i, w in enumerate(words):
        if i > 0 and i % 3 == 0:
            result.append("________")
        else:
            result.append(w)
    return " ".join(result)


# ═══════════════════════════════════════════════════════════════════════════════
#  CER WORKSHEET
# ═══════════════════════════════════════════════════════════════════════════════

def build_cer_worksheet(doc, course, unit, lesson):
    """
    Build a CER (Claim-Evidence-Reasoning) Worksheet.

    Sections:
    1. Header with unit/lesson/topic info
    2. Phenomenon or question prompt
    3. Claim section with scaffold
    4. Evidence section with data source
    5. Reasoning section with Because/But/So frame
    """
    set_document_defaults(doc)
    cer = lesson.get("cer", {})

    # 1. Header
    add_title(doc, f"{course} - CER Worksheet")
    add_styled_paragraph(
        doc,
        f"Unit {unit.get('number', '')}: {unit.get('title', '')}  |  "
        f"Lesson {lesson.get('number', '')}: {lesson.get('title', '')}",
        size=BODY_SIZE, italic=True, color=DARK_BLUE,
        alignment=WD_ALIGN_PARAGRAPH.CENTER
    )
    add_styled_paragraph(
        doc,
        f"Topic: {cer.get('topic', lesson.get('title', 'TBD'))}",
        size=HEADING_SIZE, bold=True, color=DARK_BLUE,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=Pt(4)
    )
    add_styled_paragraph(
        doc, f"Name: ________________________  Date: ____________  Period: _____",
        size=BODY_SIZE, alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=Pt(4), space_after=Pt(8)
    )
    add_section_divider(doc)

    # 2. Phenomenon / Question Prompt
    add_heading_text(doc, "Phenomenon / Question")
    p = doc.add_paragraph()
    set_paragraph_shading(p, LIGHT_BLUE)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(cer.get("phenomenon", "TBD"))
    style_run(run, size=BODY_SIZE)

    add_styled_paragraph(
        doc, "Space for visual, diagram, or data description:",
        size=SMALL_SIZE, italic=True, space_before=Pt(4)
    )
    add_lined_space(doc, num_lines=4)

    # 3. Claim Section
    add_section_divider(doc)
    add_heading_text(doc, "Claim")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    set_paragraph_shading(p, LIGHT_GREEN)
    run = p.add_run("Sentence Starter: ")
    style_run(run, size=BODY_SIZE, bold=True, color=DARK_BLUE)
    run2 = p.add_run(cer.get("claim_starter", "I claim that..."))
    style_run(run2, size=BODY_SIZE, italic=True)

    add_styled_paragraph(doc, "Write your full claim:", size=BODY_SIZE, bold=True, space_before=Pt(6))
    add_lined_space(doc, num_lines=4)

    # 4. Evidence Section
    add_section_divider(doc)
    add_heading_text(doc, "Evidence")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    set_paragraph_shading(p, LIGHT_BLUE)
    run = p.add_run("Data Source: ")
    style_run(run, size=BODY_SIZE, bold=True, color=DARK_BLUE)
    run2 = p.add_run(cer.get("evidence_source", "Use data from your investigation to support your claim."))
    style_run(run2, size=BODY_SIZE, italic=True)

    add_styled_paragraph(doc, "Cite your evidence (include specific data, observations, or measurements):",
                         size=BODY_SIZE, bold=True, space_before=Pt(6))
    add_lined_space(doc, num_lines=5)

    # 5. Reasoning Section
    add_section_divider(doc)
    add_heading_text(doc, "Reasoning")
    reasoning_frame = cer.get("reasoning_frame", "This evidence supports my claim because...")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    set_paragraph_shading(p, LIGHT_GREEN)
    run = p.add_run("Because/But/So Frame: ")
    style_run(run, size=BODY_SIZE, bold=True, color=DARK_BLUE)
    run2 = p.add_run(reasoning_frame)
    style_run(run2, size=BODY_SIZE, italic=True)

    add_styled_paragraph(doc, "Explain how your evidence supports your claim:",
                         size=BODY_SIZE, bold=True, space_before=Pt(6))
    add_lined_space(doc, num_lines=6)

    return doc


# ═══════════════════════════════════════════════════════════════════════════════
#  TEMPLATE (BLANK) VERSIONS
# ═══════════════════════════════════════════════════════════════════════════════

def build_lesson_plan_template(doc, course):
    """Build a blank lesson plan template with placeholder fields."""
    placeholder_unit = {
        "number": "___",
        "title": "[Unit Title]",
        "performance_expectations": ["[PE Code]"],
    }
    placeholder_lesson = {
        "number": "___",
        "title": "[Lesson Title]",
        "circle_prompt": "[Community prompt]",
        "do_now": "[Review/hook activity]",
        "learning_target": "At the end of 42 minutes I can [learning target].",
        "standards": {
            "pe": ["[PE]"],
            "sep": ["[SEP]"],
            "ccc": ["[CCC]"],
            "dci": ["[DCI]"],
        },
        "materials": ["[Material 1]", "[Material 2]"],
        "five_e": {
            "engage": {"time": "5 min", "description": "[Engage description]"},
            "explore": {"time": "15 min", "description": "[Explore description]"},
            "explain": {"time": "10 min", "description": "[Explain description]"},
            "elaborate": {"time": "8 min", "description": "[Elaborate description]"},
            "evaluate": {"time": "4 min", "description": "[Evaluate description]"},
        },
        "differentiation": {
            "approaching": "[Approaching accommodations]",
            "on_level": "[On-level expectations]",
            "advanced": "[Advanced extensions]",
            "ell": "[ELL supports]",
            "iep": "[IEP modifications]",
        },
        "crse_connection": "[CRSE connection]",
        "hochman_activity": "[Hochman writing activity]",
    }
    return build_lesson_plan(doc, course, placeholder_unit, placeholder_lesson)


def build_unit_plan_template(doc, course):
    """Build a blank unit plan template."""
    placeholder_unit = {
        "number": "___",
        "title": "[Unit Title]",
        "days": "___",
        "performance_expectations": ["[PE Code]"],
        "anchoring_phenomenon": "[Anchoring phenomenon description]",
        "essential_questions": ["[Essential Question 1]", "[Essential Question 2]"],
        "crse_theme": "[CRSE Theme]",
        "assessment_overview": {
            "formative": ["[Formative assessment]"],
            "summative": ["[Summative assessment]"],
        },
        "investigations": [],
        "lessons": [],
    }
    return build_unit_plan(doc, course, placeholder_unit)


def build_cer_template(doc, course):
    """Build a blank CER worksheet template."""
    placeholder_unit = {
        "number": "___",
        "title": "[Unit Title]",
    }
    placeholder_lesson = {
        "number": "___",
        "title": "[Lesson Title]",
        "cer": {
            "topic": "[Topic]",
            "phenomenon": "[Describe the phenomenon or question here]",
            "claim_starter": "[Sentence starter for claim]",
            "evidence_source": "[Describe where students should find evidence]",
            "reasoning_frame": "[Because/But/So reasoning frame]",
        },
    }
    return build_cer_worksheet(doc, course, placeholder_unit, placeholder_lesson)
