"""
Document styling constants and helpers for the Curricula DOCX generator.

Provides consistent formatting across all generated document types:
Unit Plans, Lesson Plans, Student Notes, and CER Worksheets.
"""

from docx.shared import Pt, Inches, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy


# ── Color Palette ──────────────────────────────────────────────────────────────

DARK_BLUE = RGBColor(0x1F, 0x4E, 0x79)       # #1F4E79  headers
LIGHT_BLUE = "D6E4F0"                          # accent boxes (hex string for shading)
LIGHT_GREEN = "E2EFDA"                         # accent boxes
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
TABLE_HEADER_BG = "1F4E79"                     # dark blue for table header rows
LIGHT_GRAY = "F2F2F2"


# ── Font Settings ──────────────────────────────────────────────────────────────

FONT_NAME = "Calibri"
TITLE_SIZE = Pt(14)
HEADING_SIZE = Pt(12)
BODY_SIZE = Pt(11)
SMALL_SIZE = Pt(9)

DISTRICT_NAME = "Valley Stream Central High School District"


# ── Margin Defaults ────────────────────────────────────────────────────────────

MARGIN_TOP = Inches(0.75)
MARGIN_BOTTOM = Inches(0.75)
MARGIN_LEFT = Inches(0.75)
MARGIN_RIGHT = Inches(0.75)


# ── Helper Functions ───────────────────────────────────────────────────────────

def set_document_defaults(doc):
    """Apply default margins, font, and footer to the document."""
    for section in doc.sections:
        section.top_margin = MARGIN_TOP
        section.bottom_margin = MARGIN_BOTTOM
        section.left_margin = MARGIN_LEFT
        section.right_margin = MARGIN_RIGHT

        # Add footer with district name
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(DISTRICT_NAME)
        run.font.name = FONT_NAME
        run.font.size = SMALL_SIZE
        run.font.color.rgb = DARK_BLUE


def style_run(run, font_name=FONT_NAME, size=BODY_SIZE, bold=False, italic=False,
              color=BLACK, underline=False):
    """Apply font styling to a run."""
    run.font.name = font_name
    run.font.size = size
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.underline = underline
    return run


def add_styled_paragraph(doc, text, size=BODY_SIZE, bold=False, italic=False,
                         color=BLACK, alignment=WD_ALIGN_PARAGRAPH.LEFT,
                         space_before=Pt(0), space_after=Pt(4)):
    """Add a paragraph with consistent styling."""
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = space_before
    p.paragraph_format.space_after = space_after
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_title(doc, text):
    """Add a document title in dark blue, centered, 14pt bold."""
    return add_styled_paragraph(
        doc, text,
        size=TITLE_SIZE, bold=True, color=DARK_BLUE,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=Pt(6), space_after=Pt(6)
    )


def add_heading_text(doc, text, level=1):
    """Add a heading-style paragraph (not using built-in heading styles)."""
    size = HEADING_SIZE if level == 1 else BODY_SIZE
    return add_styled_paragraph(
        doc, text,
        size=size, bold=True, color=DARK_BLUE,
        space_before=Pt(8), space_after=Pt(4)
    )


def add_shaded_box(doc, label, content, shade_color=LIGHT_BLUE):
    """Add a shaded box with a bold label and body text."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)

    # Apply shading to the paragraph
    set_paragraph_shading(p, shade_color)

    # Bold label
    run_label = p.add_run(label + ": ")
    style_run(run_label, bold=True, size=BODY_SIZE, color=DARK_BLUE)

    # Body content
    run_body = p.add_run(content)
    style_run(run_body, size=BODY_SIZE)

    return p


def set_paragraph_shading(paragraph, color_hex):
    """Apply background shading to a paragraph."""
    shading_elm = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>'
    )
    paragraph._p.get_or_add_pPr().append(shading_elm)


def set_cell_shading(cell, color_hex):
    """Apply background shading to a table cell."""
    shading_elm = parse_xml(
        f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>'
    )
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_styled_table(doc, headers, rows, col_widths=None):
    """
    Create a bordered table with a dark-blue header row.

    Parameters:
        headers: list of column header strings
        rows: list of lists (each inner list = one row of cell texts)
        col_widths: optional list of Inches values for column widths
    """
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    hdr_row = table.rows[0]
    for i, header in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        style_run(run, size=BODY_SIZE, bold=True, color=WHITE)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, TABLE_HEADER_BG)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_text))
            style_run(run, size=BODY_SIZE)
            # Alternate row shading
            if r_idx % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)

    # Column widths
    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = width

    return table


def add_bullet_list(doc, items, indent_level=0):
    """Add a bulleted list of items."""
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        if indent_level > 0:
            p.paragraph_format.left_indent = Inches(0.25 * indent_level)
        run = p.add_run(str(item))
        style_run(run, size=BODY_SIZE)


def add_lined_space(doc, num_lines=3):
    """Add blank lined space for student writing."""
    for _ in range(num_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("_" * 80)
        style_run(run, size=BODY_SIZE, color=RGBColor(0xBB, 0xBB, 0xBB))


def add_section_divider(doc):
    """Add a thin horizontal rule as a visual separator."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="1F4E79"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
