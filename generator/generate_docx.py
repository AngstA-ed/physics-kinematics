#!/usr/bin/env python3
"""
Main script for the Curricula DOCX generator.

Reads JSON course configuration files and produces formatted Word documents
for a high school science curriculum. Generates four document types:
Unit Plans, Lesson Plans, Student Notes, and CER Worksheets.

Usage:
    python -m generator.generate_docx --course chemistry
    python -m generator.generate_docx --course physics --templates-only
"""

import argparse
import json
import os
import sys
from pathlib import Path

from docx import Document

from generator.templates import (
    build_unit_plan,
    build_lesson_plan,
    build_student_notes,
    build_cer_worksheet,
    build_lesson_plan_template,
    build_unit_plan_template,
    build_cer_template,
)


BASE_DIR = Path(__file__).resolve().parent.parent
CONFIGS_DIR = Path(__file__).resolve().parent / "configs"


def load_config(course: str) -> dict:
    """Load and return the JSON config for the given course."""
    config_path = CONFIGS_DIR / f"{course}.json"
    if not config_path.exists():
        print(f"ERROR: Config file not found: {config_path}")
        sys.exit(1)
    with open(config_path, "r", encoding="utf-8") as f:
        return json.load(f)


def safe_filename(text: str) -> str:
    """Sanitize text for use as a directory or file name."""
    text = text.replace("&", "and")
    return "".join(c if c.isalnum() or c in " -_" else "" for c in text).strip()


def save_doc(doc: Document, path: Path) -> None:
    """Save a Document to disk, creating parent dirs as needed."""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def generate_templates(course_name: str, output_root: Path) -> None:
    """Generate blank template documents."""
    templates_dir = output_root / "Templates"
    templates_dir.mkdir(parents=True, exist_ok=True)

    # Lesson Plan Template
    print(f"  -> Lesson Plan Template")
    doc = Document()
    build_lesson_plan_template(doc, course_name)
    save_doc(doc, templates_dir / "Lesson Plan Template.docx")

    # Unit Plan Template
    print(f"  -> Unit Plan Template")
    doc = Document()
    build_unit_plan_template(doc, course_name)
    save_doc(doc, templates_dir / "Unit Plan Template.docx")

    # CER Template
    print(f"  -> CER Worksheet Template")
    doc = Document()
    build_cer_template(doc, course_name)
    save_doc(doc, templates_dir / "CER Worksheet Template.docx")


def generate_course(config: dict) -> None:
    """Generate all documents for a course based on its config."""
    course_name = config.get("course", "Course")
    output_root = BASE_DIR / course_name

    print(f"\nGenerating documents for: {course_name}")
    print(f"Output directory: {output_root}\n")

    # Templates
    print("Creating templates...")
    generate_templates(course_name, output_root)

    # Units
    units = config.get("units", [])
    if not units:
        print("  No units defined in config.")
        return

    for unit in units:
        unit_num = unit.get("number", 0)
        unit_title = unit.get("title", "Untitled")
        unit_dir_name = f"{unit_num:02d}-{safe_filename(unit_title)}"
        unit_dir = output_root / "Units" / unit_dir_name

        print(f"\nUnit {unit_num}: {unit_title}")

        # Unit Plan
        print(f"  -> Unit Plan")
        doc = Document()
        build_unit_plan(doc, course_name, unit)
        save_doc(doc, unit_dir / "Unit Plan.docx")

        # Lessons
        lessons = unit.get("lessons", [])
        for lesson in lessons:
            lesson_num = lesson.get("number", 0)
            lesson_title = lesson.get("title", "Untitled")
            lesson_dir_name = f"Lesson {lesson_num} - {safe_filename(lesson_title)}"
            lesson_dir = unit_dir / lesson_dir_name

            print(f"  Lesson {lesson_num}: {lesson_title}")

            # Lesson Plan
            print(f"    -> Lesson Plan")
            doc = Document()
            build_lesson_plan(doc, course_name, unit, lesson)
            save_doc(doc, lesson_dir / "Lesson Plan.docx")

            # Student Notes
            print(f"    -> Student Notes")
            doc = Document()
            build_student_notes(doc, course_name, unit, lesson)
            save_doc(doc, lesson_dir / "Student Notes.docx")

            # CER Worksheet (only if lesson has CER data)
            cer = lesson.get("cer")
            if cer:
                print(f"    -> CER Worksheet")
                doc = Document()
                build_cer_worksheet(doc, course_name, unit, lesson)
                save_doc(doc, lesson_dir / "CER Worksheet.docx")
            else:
                print(f"    -- CER Worksheet skipped (no CER data)")

    print(f"\nDone! All documents saved under: {output_root}")


def main():
    parser = argparse.ArgumentParser(
        description="Generate DOCX curriculum documents from JSON configuration."
    )
    parser.add_argument(
        "--course",
        required=True,
        help="Course name (matches a JSON file in generator/configs/, e.g., chemistry, physics, test)"
    )
    parser.add_argument(
        "--templates-only",
        action="store_true",
        help="Generate only blank template documents, not full unit/lesson documents"
    )
    args = parser.parse_args()

    config = load_config(args.course)
    course_name = config.get("course", args.course.title())

    if args.templates_only:
        output_root = BASE_DIR / course_name
        print(f"\nGenerating templates only for: {course_name}")
        generate_templates(course_name, output_root)
        print("\nDone!")
    else:
        generate_course(config)


if __name__ == "__main__":
    main()
